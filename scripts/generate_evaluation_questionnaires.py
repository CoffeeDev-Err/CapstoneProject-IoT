from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

BLUE = "000000"
PALE_BLUE = "F2F2F2"
SLATE = "111111"
WHITE = "FFFFFF"


CRITERIA = [
    (
        "A. FUNCTIONAL SUITABILITY",
        "Coverage: Functional Completeness (FCO), Functional Correctness (FCR), and Functional Appropriateness (FAP)",
    ),
    (
        "B. USABILITY",
        "Coverage: Learnability (L), Operability (O), User Error Protection (UEP), and User Interface Aesthetics (UIA)",
    ),
    (
        "C. RELIABILITY",
        "Coverage: Availability (AV), Fault Tolerance (FT), and Recoverability (RC)",
    ),
    (
        "D. PERFORMANCE EFFICIENCY",
        "Coverage: Time Behavior (TB), Resource Utilization (RU), and Capacity (C)",
    ),
    (
        "E. SECURITY",
        "Coverage: Confidentiality (CF), Integrity (IN), Accountability (AC), and Authenticity (AU)",
    ),
]


MOBILE_ITEMS = [
    [
        "The mobile application provides the functions needed to view my current deployment, assigned tasks, and duty status. (FCO)",
        "The mobile application provides the functions needed for location updating, backup requests, notifications, and report submission. (FCO)",
        "The mobile application provides access to submitted-report history and account functions needed during duty. (FCO)",
        "Deployment, task, and schedule details displayed in the application match my assigned duty information. (FCR)",
        "The application displays the correct latest GPS status and location information received for my account. (FCR)",
        "Notifications and report history display the correct event, status, and related record. (FCR)",
        "The map, task, and notification functions help me understand and perform my assigned duties effectively. (FAP)",
        "The backup-request and reporting functions help me request assistance and document incidents effectively. (FAP)",
    ],
    [
        "The sign-in and verification-code process is easy to learn. (L)",
        "The purpose of the Map, Tasks, Reports, and Account screens is easy to understand. (L)",
        "I can move between the main screens without difficulty. (O)",
        "Buttons, filters, forms, and map controls are easy to operate. (O)",
        "Validation messages clearly explain missing or incorrect information. (UEP)",
        "Confirmations, saved drafts, and retry options help prevent lost, accidental, or duplicate actions. (UEP)",
        "Text, labels, buttons, and status information are easy to read. (UIA)",
        "Colors, icons, spacing, and screen layouts provide a consistent and professional appearance. (UIA)",
    ],
    [
        "The mobile application is accessible and ready during normal duty operations. (AV)",
        "My assigned deployments, tasks, and submitted-report history are available when needed. (AV)",
        "GPS updates, notifications, backup requests, and reporting functions work consistently during normal use. (AV)",
        "The application provides clear and safe behavior when the internet or mobile signal is temporarily unavailable. (FT)",
        "Previously viewed tasks and submitted reports remain available when temporary connectivity problems occur. (FT)",
        "The application avoids unnecessary duplicate report submissions when offline synchronization is retried. (FT)",
        "The application returns to normal operation after internet connectivity is restored. (RC)",
        "Pending reports or saved report drafts can be recovered and synchronized after connectivity returns. (RC)",
    ],
    [
        "Mobile screens load within an acceptable amount of time. (TB)",
        "Sign-in, search, filtering, backup-request, and report-submission actions respond promptly. (TB)",
        "New notifications and GPS information appear soon after the related data are received. (TB)",
        "The application operates without causing noticeable device slowdown during normal use. (RU)",
        "Map movement, scrolling, card expansion, and screen transitions operate smoothly. (RU)",
        "The application performs ordinary tasks without unnecessary repeated loading or excessive resource use. (RU)",
        "The application remains responsive while displaying multiple map markers, tasks, reports, or notifications. (C)",
        "The application remains responsive during repeated normal operations throughout a duty period. (C)",
    ],
    [
        "The application limits personnel, location, and operational information to an authenticated account. (CF)",
        "I can access only the information and functions permitted for my police-personnel role. (CF)",
        "Input validation and confirmations help prevent invalid or accidental changes to important information. (IN)",
        "Submitted reports, evidence, and report history remain complete and accurate after submission or correction. (IN)",
        "My submitted reports, backup requests, and important actions are associated with my authenticated account. (AC)",
        "Notifications and history information provide the related event, time, and status needed to trace important actions. (AC)",
        "The Login ID, password, and verification-code process helps confirm that only the authorized user can sign in. (AU)",
        "Logout and session controls prevent continued access after I sign out or when my session expires. (AU)",
    ],
]


WEB_ITEMS = [
    [
        "The web portal provides the map and location-monitoring functions needed to supervise police personnel. (FCO)",
        "The web portal provides the functions needed to manage personnel accounts, deployments, tasks, and backup requests. (FCO)",
        "The web portal provides report validation, correction, notification, report-generation, and analytics functions. (FCO)",
        "Personnel identity, duty status, deployment, and latest received GPS information are displayed correctly. (FCR)",
        "Deployment, task, backup-request, and report details are recorded and displayed correctly. (FCR)",
        "Report validation, correction reason, reviewer, and history information are recorded correctly. (FCR)",
        "Map, search, and filtering functions help me locate personnel and monitor operations effectively. (FAP)",
        "Personnel management, deployment, reporting, and analytics functions support supervision and decision-making. (FAP)",
    ],
    [
        "The sign-in and verification-code process is easy to learn. (L)",
        "The purpose of the dashboard, map, navigation menu, tables, forms, and icons is easy to understand. (L)",
        "I can move between monitoring and management functions without difficulty. (O)",
        "Search, filter, map, table, form, and report controls are easy to operate. (O)",
        "Validation messages clearly explain missing or incorrect information. (UEP)",
        "Confirmations, restrictions, and retry options help prevent accidental, invalid, or duplicate actions. (UEP)",
        "Text, tables, charts, labels, buttons, and status information are easy to read. (UIA)",
        "Colors, icons, spacing, and page layouts provide a consistent and professional appearance. (UIA)",
    ],
    [
        "The web portal is accessible and ready when required for normal supervision. (AV)",
        "The dashboard, map, personnel, deployment, task, report, and analytics pages are available when needed. (AV)",
        "Live location updates, notifications, search, and management functions work consistently during normal use. (AV)",
        "The portal provides clear and safe behavior when the internet connection is temporarily interrupted. (FT)",
        "Delayed, stale, or temporarily unavailable GPS readings are clearly identified without making the portal unusable. (FT)",
        "Failed requests provide an appropriate retry option without creating unnecessary duplicate records. (FT)",
        "The portal returns to normal operation after internet connectivity is restored. (RC)",
        "Previously confirmed operational information reloads correctly after refreshing or reopening the portal. (RC)",
    ],
    [
        "Web pages and dashboard information load within an acceptable amount of time. (TB)",
        "Search, filtering, assignment, saving, validation, and report-generation actions respond promptly. (TB)",
        "New GPS readings, map-marker changes, and notifications appear soon after the data are received. (TB)",
        "The portal operates smoothly in the intended web browser during normal use. (RU)",
        "Map movement, marker clustering, tables, charts, and page navigation operate without excessive lag. (RU)",
        "The portal performs ordinary operations without unnecessary repeated loading or excessive resource use. (RU)",
        "The map remains responsive while displaying the normal number of personnel markers and operational areas. (C)",
        "The portal remains responsive while displaying multiple personnel, deployments, tasks, reports, and notifications. (C)",
    ],
    [
        "The portal limits personnel locations, reports, evidence, and operational information to authorized accounts. (CF)",
        "Users can access only the functions permitted for their assigned supervisor role. (CF)",
        "The portal prevents invalid or unauthorized changes to personnel, deployment, task, and report records. (IN)",
        "Validated reports, correction reasons, and revision history remain complete and accurate after updates. (IN)",
        "Account, deployment, task, and report actions are associated with the responsible authenticated account. (AC)",
        "Report reviews and other important actions retain the responsible user, action, status, and time. (AC)",
        "The Login ID, password, and verification-code process helps confirm the identity of the authorized user. (AU)",
        "Session and logout controls keep access associated with the correct account and end access when required. (AU)",
    ],
]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, size: float = 8.5, bold: bool = False,
                  color: str = SLATE, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_layout(table, widths: list[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_layout = table_properties.find(qn("w:tblLayout"))
    if table_layout is None:
        table_layout = OxmlElement("w:tblLayout")
        table_properties.append(table_layout)
    table_layout.set(qn("w:type"), "fixed")
    if widths:
        table_width = table_properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_properties.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(round(sum(widths) * 1440)))
        grid_columns = table._tbl.tblGrid.gridCol_lst
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
            if index < len(grid_columns):
                grid_columns[index].set(qn("w:w"), str(round(width * 1440)))
    for row in table.rows:
        prevent_row_split(row)


def add_title(document: Document, platform_title: str, audience: str) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("GEOSENTRI USER EVALUATION QUESTIONNAIRE")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    run = subtitle.add_run(platform_title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    project = document.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project.paragraph_format.space_after = Pt(8)
    run = project.add_run("GeoSentri: An IoT-Based Real-Time GPS Monitoring System for Police Personnel")
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(SLATE)

    info = document.add_table(rows=2, cols=2)
    set_repeat_table_layout(info, [4.45, 3.05])
    set_cell_text(info.cell(0, 0), "Respondent Code/Name (Optional): ______________________________", size=8)
    set_cell_text(info.cell(0, 1), "Date: ____________________", size=8)
    set_cell_text(info.cell(1, 0), f"Respondent Role: {audience}", size=8, bold=True)
    set_cell_text(info.cell(1, 1), "Experience Using GeoSentri: ____________________", size=8)
    info.style = "Table Grid"
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_intro(document: Document, platform_name: str, audience: str) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run("Evaluation Objective")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(
        f"Evaluate the GeoSentri {platform_name} in terms of Functional Suitability, Usability, "
        "Reliability, Performance Efficiency, and Security based on ISO/IEC 25010:2011."
    )
    run.font.name = "Arial"
    run.font.size = Pt(8.5)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run("Instructions")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(
        f"This questionnaire is intended only for {audience}. After using GeoSentri, check one box "
        "that best represents your level of agreement with each statement. Answer all items based "
        "on your actual experience. Each of the five quality criteria contains eight statements, "
        "for a total of 40 statements."
    )
    run.font.name = "Arial"
    run.font.size = Pt(8.5)

    scale = document.add_table(rows=2, cols=5)
    scale.style = "Table Grid"
    set_repeat_table_layout(scale, [1.52, 1.52, 1.52, 1.52, 1.52])
    labels = [("5", "Strongly Agree"), ("4", "Agree"), ("3", "Neutral"), ("2", "Disagree"), ("1", "Strongly Disagree")]
    for index, (score, meaning) in enumerate(labels):
        set_cell_text(scale.cell(0, index), score, size=8.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(scale.cell(0, index), BLUE)
        set_cell_text(scale.cell(1, index), meaning, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_criterion(document: Document, title: str, coverage: str, items: list[str]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(1)
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    description = document.add_paragraph()
    description.paragraph_format.space_after = Pt(3)
    description.paragraph_format.keep_with_next = True
    run = description.add_run(coverage)
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(SLATE)

    table = document.add_table(rows=2 + len(items), cols=7)
    table.style = "Table Grid"
    widths = [0.36, 5.39, 0.37, 0.37, 0.37, 0.37, 0.37]
    set_repeat_table_layout(table, widths)

    merged_no = table.cell(0, 0).merge(table.cell(1, 0))
    merged_statement = table.cell(0, 1).merge(table.cell(1, 1))
    merged_rating = table.cell(0, 2).merge(table.cell(0, 6))
    set_cell_text(merged_no, "NO.", size=7.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(merged_statement, "STATEMENTS", size=7.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(merged_rating, "LEVEL OF AGREEMENT", size=7.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in [merged_no, merged_statement, merged_rating]:
        shade(cell, BLUE)
    for index, score in enumerate(["5", "4", "3", "2", "1"], start=2):
        set_cell_text(table.cell(1, index), score, size=7.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.cell(1, index), BLUE)
    repeat_table_header(table.rows[0])
    repeat_table_header(table.rows[1])

    for item_number, statement in enumerate(items, start=1):
        row_index = item_number + 1
        if item_number % 2 == 0:
            for cell in table.rows[row_index].cells:
                shade(cell, PALE_BLUE)
        set_cell_text(table.cell(row_index, 0), str(item_number), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(row_index, 1), statement, size=8)
        for column in range(2, 7):
            set_cell_text(table.cell(row_index, column), "☐", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_comments(document: Document) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run("Comments or Suggestions")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    for _ in range(3):
        paragraph = document.add_paragraph("_" * 112)
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)


def add_footer(document: Document, label: str) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"GeoSentri {label} Evaluation Questionnaire | ISO/IEC 25010:2011")
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("64748B")


def build_questionnaire(filename: str, platform_title: str, platform_name: str,
                        audience: str, items: list[list[str]]) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(13)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)

    add_title(document, platform_title, audience)
    add_intro(document, platform_name, audience)
    for (title, coverage), criterion_items in zip(CRITERIA, items, strict=True):
        add_criterion(document, title, coverage, criterion_items)
    add_comments(document)
    add_footer(document, platform_title)

    output = DOCS / filename
    document.save(output)
    return output


if __name__ == "__main__":
    DOCS.mkdir(parents=True, exist_ok=True)
    outputs = [
        build_questionnaire(
            "GeoSentri_Mobile_Police_Personnel_Evaluation_Questionnaire.docx",
            "Mobile Application — Police Personnel",
            "mobile application",
            "Police Personnel",
            MOBILE_ITEMS,
        ),
        build_questionnaire(
            "GeoSentri_Web_Supervisor_Evaluation_Questionnaire.docx",
            "Web Monitoring Portal — Supervisor",
            "web monitoring portal",
            "Supervisor/Administrator",
            WEB_ITEMS,
        ),
    ]
    for output in outputs:
        print(output)
